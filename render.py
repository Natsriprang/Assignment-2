"""Render a clause list to a .docx."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from renumber import renumber_and_render


def render_docx(clauses, out_path):
    rendered = renumber_and_render(clauses)

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Construction and Installation Service Agreement")
    r.bold = True
    r.font.size = Pt(14)
    doc.add_paragraph()

    for number, clause, item_lines in rendered:
        h = doc.add_paragraph()
        r = h.add_run(f"{number}.\t{clause.title}")
        r.bold = True

        if clause.plain_paragraphs:
            for p in clause.plain_paragraphs:
                doc.add_paragraph(p)
        for line in item_lines:
            p = doc.add_paragraph()
            if line.startswith("    "):
                p.paragraph_format.left_indent = Cm(1.0)
                p.add_run(line.strip())
            else:
                p.paragraph_format.left_indent = Cm(0.5)
                p.add_run(line)

    doc.save(out_path)
    return rendered