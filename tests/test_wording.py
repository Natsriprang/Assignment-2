"""Tests for wording changes: finding the sentence, checking a rewrite, applying accepted changes, before/after markup.
Run: python -m tests.test_wording"""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document

from core.clause_data import build_library
from core.draft_state import library_for
from core.header import build_header
from core.render import render_docx
from core.renumber import renumber_and_render
from core.wording import (apply_edits, check_rewrite, core_text, diff_markup, figure_mismatches, figure_swap, find_unit, norm, removable,
                          removal_of, split_point, splice_fragment, unit_hits, units, without_existing)

LD_QUOTE = "at the rate of 0.05 (zero point zero five) percent of the Contract Price per day"
REQ = "Reduce the delay damages rate to 0.03% per day. reduce delay damages to 0.03% per day"


def clause(key, library=None):
    return next(c for c in (library or library_for({})) if c.key == key)


def old_ld():
    return find_unit(clause("completion_deadline"), LD_QUOTE)[1]


def new_ld(figure="0.03 (zero point zero three)"):
    return old_ld().replace("0.05 (zero point zero five)", figure)


# ---------------------------------------------------------------- finding the sentence
def test_the_sentence_containing_a_quote_is_found_whatever_kind_of_text_it_is():
    loc, text = find_unit(clause("completion_deadline"), LD_QUOTE)
    assert loc[0] == "item" and text.startswith("Unless otherwise agreed by the Parties") and "0.05 (zero point zero five)" in text
    assert find_unit(clause("price_payment"), "mounting system")[0][0] == "cell"
    assert find_unit(clause("notice"), "Any communication with the addressee")[0][0] == "paragraph"
    assert find_unit(clause("completion_deadline"), "  AT THE RATE of 0.05   (zero point zero five) PERCENT ")[0] == loc


def test_a_quote_in_no_sentence_or_in_several_finds_nothing():
    assert find_unit(clause("completion_deadline"), "a provision that is not there") is None
    assert find_unit(clause("completion_deadline"), "") is None
    assert find_unit(clause("price_payment"), "Payable upon completion of installation") is None          # in more than one table cell


def test_every_piece_of_text_of_a_clause_is_a_unit():
    c = clause("price_payment")
    kinds = {loc[0] for loc, _ in units(c)}
    assert kinds == {"item", "cell"} and len([1 for loc, _ in units(c) if loc[0] == "cell"]) == 5 * 4


# ---------------------------------------------------------------- checking a rewrite
def test_a_change_of_the_figure_and_its_words_is_accepted():
    assert check_rewrite(old_ld(), new_ld(), "modify", LD_QUOTE, REQ) == ""


def test_a_figure_and_its_words_must_agree():
    assert "do not agree" in check_rewrite(old_ld(), new_ld("0.03 (zero point zero five)"), "modify", LD_QUOTE, REQ)
    assert figure_mismatches("pay within 15 (fifteen) days and 7 (eight) days") == ["7 (eight)"]
    assert figure_mismatches("5% (five percent) and 0.03 (zero point zero three) percent and 120 (one hundred and twenty) days") == []
    assert figure_mismatches("15 (fourteen) days", known="15 (fourteen) days") == []            # an old mismatch is not this change's


def test_only_figures_the_user_gave_may_come_in():
    assert "figure that you did not give" in check_rewrite(old_ld(), new_ld("0.04 (zero point zero four)"), "modify", LD_QUOTE, REQ)
    assert "figure that you did not give" in check_rewrite("", "The warranty period is 36 (thirty-six) months.", "add", "", "add a 24 months warranty")
    assert check_rewrite("", "The warranty period is 24 (twenty-four) months.", "add", "", "add a 24 months warranty") == ""
    assert check_rewrite("", "Repair within 120 (one hundred and twenty) days.", "add", "", "repair within 120 days") == ""


def test_the_figure_asked_for_must_be_in_the_new_wording():
    assert "does not contain the figure" in check_rewrite("Pay within 30 (thirty) days.", "Pay within 30 (thirty) days promptly.", "modify", "30 (thirty) days", "make it 15 days")


def test_nothing_changing_or_everything_changing_is_refused():
    assert check_rewrite(old_ld(), old_ld(), "modify", LD_QUOTE, REQ) == "The wording would not change."
    rewritten = "The Owner may claim 0.03 (zero point zero three) percent of the Contract Price per day."
    assert "much more of the sentence" in check_rewrite(old_ld(), rewritten, "modify", LD_QUOTE, REQ)
    assert check_rewrite(old_ld(), "", "modify", LD_QUOTE, REQ) == "The AI did not write the new wording."
    assert "too long" in check_rewrite("A b.", "word " * 200, "add", "", "")


def test_a_change_outside_the_quoted_words_is_refused():
    elsewhere = new_ld().replace("Unless otherwise agreed by the Parties", "Unless otherwise agreed in writing by the Parties")
    assert "parts of the sentence you did not ask about" in check_rewrite(old_ld(), elsewhere, "modify", LD_QUOTE, REQ)
    near = new_ld().replace("per day", "each day")                            # inside the quote
    assert check_rewrite(old_ld(), near, "modify", LD_QUOTE, REQ) == ""


def test_removing_part_of_a_sentence():
    old = "The Contractor shall provide the Owner with weekly progress reports and photographs of the Works."
    quote = "and photographs of the Works"
    assert check_rewrite(old, "The Contractor shall provide the Owner with weekly progress reports.", "remove_part", quote, "remove the photographs") == ""
    assert "still in the new wording" in check_rewrite(old, old.replace("progress", "interim progress"), "remove_part", quote, "remove the photographs")
    assert "much more" in check_rewrite(old, "The Contractor shall provide.", "remove_part", quote, "remove the photographs")


# ---------------------------------------------------------------- repairing a proposal
LONG_QUOTE = ("the Owner has the right to claim liquidated damages from the Contractor at the rate of 0.05 (zero point zero five) percent of "
              "the Contract Price per day but not more than 5 (five) percent of the Contract Price.")
REQUEST_2 = "Reduce the liquidated damages rate to 0.03% reduce the delay damage to 0.03%"


def test_only_the_quoted_words_given_back_are_put_into_the_sentence():
    old = find_unit(clause("completion_deadline"), LD_QUOTE)[1]
    fragment = LONG_QUOTE.replace("0.05 (zero point zero five)", "0.03 (zero point zero three)")
    assert "much more of the sentence" in check_rewrite(old, fragment, "modify", LONG_QUOTE, REQUEST_2)           # the reported refusal
    whole = splice_fragment(old, LONG_QUOTE, fragment)
    assert whole == old.replace("0.05 (zero point zero five)", "0.03 (zero point zero three)")
    assert check_rewrite(old, whole, "modify", LONG_QUOTE, REQUEST_2) == ""


def test_a_fragment_in_the_middle_keeps_the_words_on_both_sides():
    old = "The Owner shall pay the Contractor within 30 (thirty) working days after receipt of the invoice, without any deduction."
    got = splice_fragment(old, "within 30 (thirty) working days", "within 15 (fifteen) working days")
    assert got == "The Owner shall pay the Contractor within 15 (fifteen) working days after receipt of the invoice, without any deduction."
    ended = splice_fragment("Pay within 30 (thirty) days.", "within 30 (thirty) days", "within 15 (fifteen) days.")
    assert ended == "Pay within 15 (fifteen) days."                                           # a full stop is not doubled


def test_a_whole_sentence_or_an_unrelated_text_is_not_treated_as_a_fragment():
    old = find_unit(clause("completion_deadline"), LD_QUOTE)[1]
    assert splice_fragment(old, LD_QUOTE, old.replace("0.05 (zero point zero five)", "0.03 (zero point zero three)")) == ""
    assert splice_fragment(old, LD_QUOTE, "The Contractor shall do something entirely different.") == ""
    assert splice_fragment(old, "a quote that is not there", "x") == "" and splice_fragment(old, LD_QUOTE, "") == ""


def test_the_system_writes_one_figure_and_its_words_itself():
    old = find_unit(clause("completion_deadline"), LD_QUOTE)[1]
    got = figure_swap(old, LD_QUOTE, "reduce the rate to 0.03% per day")
    assert got == old.replace("0.05 (zero point zero five)", "0.03 (zero point zero three)")        # one figure in the quote: replaced
    one = "at the rate of 0.05 (zero point zero five) percent of the Contract Price per day"
    got = figure_swap(old, one, "reduce the rate to 0.03% per day")
    assert got == old.replace("0.05 (zero point zero five)", "0.03 (zero point zero three)")
    assert check_rewrite(old, got, "modify", one, "reduce the rate to 0.03% per day") == ""
    pay = "The Owner shall pay within 30 (thirty) working days."
    assert figure_swap(pay, "within 30 (thirty) working days", "make it 45 working days") == "The Owner shall pay within 45 (forty-five) working days."
    assert figure_swap(pay, "within 30 (thirty) working days", "make it 120 working days") == "The Owner shall pay within 120 (one hundred and twenty) working days."
    pct = "not more than 5% (five percent) of the Contract Price"
    assert figure_swap(pct, "5% (five percent)", "cap it at 8%") == "not more than 8% (eight percent) of the Contract Price"
    assert figure_swap(pct, "5% (five percent)", "cap it at 2.5%") == "not more than 2.5% (two point five percent) of the Contract Price"


def test_the_system_does_not_guess_which_figure_or_what_the_new_one_is():
    old = find_unit(clause("completion_deadline"), LD_QUOTE)[1]
    assert figure_swap(old, LONG_QUOTE, "reduce the rate to 0.03% per day") == ""                    # two figures in the quote: which one?
    one = "at the rate of 0.05 (zero point zero five) percent of the Contract Price per day"
    assert figure_swap(old, one, "reduce the rate") == ""                                            # no new figure
    assert figure_swap(old, one, "reduce the rate to 0.03% per day or 0.02%") == ""                 # two new figures
    assert figure_swap(old, one, "reduce the rate to 0.0123% per day") == ""                         # more than two decimals
    assert figure_swap(old, "a quote that is not there", "to 0.03%") == ""
    assert figure_swap("Pay within thirty days.", "within thirty days", "make it 15 days") == ""      # no digits-and-words pair to replace


# ---------------------------------------------------------------- taking a whole point or sentence out
def take_out(key, quote):
    loc, unit = find_unit(clause(key), quote)
    edit = {"id": f"rm-{key}-{len(unit)}", "clause_key": key, "change_type": "remove", "old": unit, "new": "", "summary": "s"}
    return clause(key, library_for({}, edits=[edit])), loc


def test_a_quote_that_is_all_of_a_unit_removes_the_unit():
    point = "7) epidemic."
    assert removal_of(point, "epidemic") and removal_of(point, "7) epidemic.") and removal_of(point, "epidemic.")
    joined = "6) any change in law or government action; or"
    assert removal_of(joined, "any change in law or government action") and removal_of(joined, joined) and removal_of("(a) x, and", "x")
    assert not removal_of("6) any change in law or government action or the Owner's consent.", "any change in law or government action")
    assert not removal_of(point, "epid") and not removal_of(point, "something else")
    sentence = "The Owner shall pay the Contractor within 30 (thirty) working days."
    assert removal_of(sentence, sentence) and not removal_of(sentence, "within 30 (thirty) working days")
    assert removable(("lettered", 1, 6)) and removable(("item", 0)) and removable(("paragraph", 1)) and not removable(("cell", 0, 1, 3))


def test_the_last_point_of_a_list_goes_and_the_one_before_it_takes_the_full_stop():
    c, loc = take_out("force_majeure", "epidemic")
    points = c.items[1].lettered
    assert loc == ("lettered", 1, 6) and len(points) == 6 and points[-1] == "6) any change in law or government action."
    assert all(p.endswith("; or") for p in points[:-1]) and not any("epidemic" in p for p in points)


def test_a_middle_point_goes_and_the_labels_are_counted_again():
    c, _ = take_out("force_majeure", "protests and/or other strikes")
    points = c.items[1].lettered
    assert [p[:2] for p in points] == ["1)", "2)", "3)", "4)", "5)", "6)"] and points[2].startswith("3) war, invasion")
    assert points[-1] == "6) epidemic." and all(p.endswith("; or") for p in points[:-1])


def test_the_or_goes_back_one_when_the_point_that_carried_it_is_taken_out():
    c, _ = take_out("confidentiality", "Confidential information which was already in")
    first, last = c.items[2].lettered
    assert first.startswith("(1)") and first.endswith("; or") and last.startswith("(2)") and last.endswith(".")


def test_lettered_lists_are_counted_in_letters_and_the_last_point_keeps_the_full_stop():
    c, _ = take_out("termination", "abandons the Works")
    a, b = c.items[1].lettered
    assert a.startswith("(a)") and a.endswith(";") and b.startswith("(b)") and b.endswith(".")
    c, _ = take_out("termination", "fails to comply with the notice under clause 1")
    assert [p[:3] for p in c.items[1].lettered] == ["(a)", "(b)"] and c.items[1].lettered[-1].endswith("reasonable cause.")


def test_the_only_point_of_a_list_cannot_be_taken_out_on_its_own():
    library = library_for({})
    c = clause("force_majeure", library)
    c.items[1].lettered = ["1) epidemic."]
    edit = {"id": "x", "clause_key": "force_majeure", "change_type": "remove", "old": "1) epidemic.", "new": "", "summary": "s"}
    failed = []
    apply_edits(library, [edit], failed)
    assert failed == [edit] and c.items[1].lettered == ["1) epidemic."]


def test_a_numbered_sentence_and_a_paragraph_can_be_taken_out():
    notice = clause("notice")
    second = notice.plain_paragraphs[1]
    c = clause("notice", library_for({}, edits=[{"id": "p", "clause_key": "notice", "change_type": "remove", "old": second, "new": "", "summary": "s"}]))
    assert c.plain_paragraphs == [notice.plain_paragraphs[0]]
    dead = clause("completion_deadline").items[1].text
    c = clause("completion_deadline", library_for({}, edits=[{"id": "i", "clause_key": "completion_deadline", "change_type": "remove", "old": dead, "new": "", "summary": "s"}]))
    assert len(c.items) == len(clause("completion_deadline").items) - 1 and not any(norm(i.text) == norm(dead) for i in c.items)
    lines = next(items for n, cl, items in renumber_and_render(library_for({}, edits=[{"id": "i", "clause_key": "completion_deadline", "change_type": "remove", "old": dead, "new": "", "summary": "s"}])) if cl.key == "completion_deadline")
    assert [l.split(" ", 1)[0] for l in lines] == [f"6.{n}" for n in range(1, len(lines) + 1)]            # the numbers close up by themselves


def test_a_removal_whose_sentence_is_gone_is_not_applied_and_the_rest_of_the_list_is_untouched():
    failed = []
    stale = {"id": "s", "clause_key": "force_majeure", "change_type": "remove", "old": "9) not a point.", "new": "", "summary": "s"}
    library_for({}, edits=[stale], failed=failed)
    assert failed == [stale]


def test_a_removal_reaches_the_word_file():
    edit = {"id": "r", "clause_key": "force_majeure", "change_type": "remove", "old": "7) epidemic.", "new": "", "summary": "s"}
    buf = io.BytesIO()
    render_docx(library_for({}, edits=[edit]), buf, header=build_header(None))
    buf.seek(0)
    text = "\n".join(p.text for p in Document(buf).paragraphs)
    assert "epidemic" not in text and "6) any change in law or government action." in text


def test_the_before_and_after_of_a_removal():
    before, after = diff_markup("7) epidemic.", "")
    assert before == "**7) epidemic.**" and after == ""


# ---------------------------------------------------------------- a quote that is a piece of the contract is found
def test_punctuation_and_quotation_marks_at_the_ends_of_a_quote_do_not_matter():
    c = clause("force_majeure")
    want = find_unit(c, "any change in law or government action")[0]
    for quote in ("any change in law or government action.", "any change in law or government action;", "\u201cany change in law or government action.\u201d",
                  "  Any Change in Law or Government Action ;", "any change in law or government action; or"):
        assert find_unit(c, quote) is not None and find_unit(c, quote)[0] == want, quote


def test_a_quote_in_a_sentence_and_in_its_lettered_point_is_the_point():
    library = library_for({})
    c = clause("force_majeure", library)
    point = c.items[1].lettered[5]
    c.items[1].text += " " + point                                  # the sentence now repeats the point
    assert unit_hits(c, "any change in law or government action") == 2
    assert find_unit(c, "any change in law or government action")[0] == ("lettered", 1, 5), "the more specific unit wins"


def test_a_quote_in_two_equally_specific_units_is_not_guessed():
    c = clause("force_majeure", library_for({}))
    c.items[1].lettered.append(c.items[1].lettered[5])
    assert unit_hits(c, "any change in law or government action") == 2 and find_unit(c, "any change in law or government action") is None
    assert find_unit(c, "nothing like this at all") is None and unit_hits(c, "nothing like this at all") == 0 and find_unit(c, "  ") is None


# ---------------------------------------------------------------- the new wording is bold only where it differs
def test_a_new_wording_that_is_mostly_new_is_shown_plain_and_a_small_change_is_not():
    old = "7) epidemic."
    new = "7) epidemic, other than for the payment obligations of the Owner."
    before, after = diff_markup(old, new)
    assert after == "7) epidemic, other than for the payment obligations of the Owner." and "**" not in after, "mostly new: plain"
    assert before == "7) **epidemic.**", "the words to be changed are still bold in the current wording"
    sentence = "The Owner shall pay the Contractor within 30 (thirty) working days after receipt of the invoice."
    before, after = diff_markup(sentence, sentence.replace("30 (thirty)", "15 (fifteen)"))
    assert "**15 (fifteen)**" in after and "**30 (thirty)**" in before and after.count("**") == 2, "a small change is marked"
    assert diff_markup("same", "same") == ("same", "same") and diff_markup("a b", "")[0] == "**a b**" and diff_markup("", "x y")[1] == "x y"


# ---------------------------------------------------------------- changes to neighbouring points do not depend on their order
def fm_edits():
    c = clause("force_majeure")
    p6 = c.items[1].lettered[5]
    return ({"id": "e7", "clause_key": "force_majeure", "change_type": "remove", "old": "7) epidemic.", "new": "", "summary": "s"},
            {"id": "e6", "clause_key": "force_majeure", "change_type": "remove", "old": p6, "new": "", "summary": "s"})


def points(edits):
    failed = []
    lib = library_for({}, edits=list(edits), failed=failed)
    return [p for p in clause("force_majeure", lib).items[1].lettered], [f["id"] for f in failed]


def test_two_points_of_one_list_can_be_taken_out_together_in_either_order():
    e7, e6 = fm_edits()
    a, failed_a = points([e7, e6])
    b, failed_b = points([e6, e7])
    assert a == b and failed_a == [] and failed_b == []
    assert [p[:2] for p in a] == ["1)", "2)", "3)", "4)", "5)"] and a[-1].endswith("government intervention.") and all(p.endswith("; or") for p in a[:-1])


def test_taking_one_of_them_back_leaves_the_other_working():
    e7, e6 = fm_edits()
    only6, failed = points([e6])
    assert failed == [] and only6[-1].startswith("6) epidemic.") and len(only6) == 6
    only7, failed = points([e7])
    assert failed == [] and only7[-1].startswith("6) any change in law or government action.") and len(only7) == 6


def test_a_wording_change_to_a_point_and_the_removal_of_its_neighbour_work_in_either_order():
    e7, _ = fm_edits()
    c = clause("force_majeure")
    p3 = c.items[1].lettered[2]
    modify = {"id": "m3", "clause_key": "force_majeure", "change_type": "modify", "old": p3, "new": p3.replace("protests and/or other strikes", "protests and/or any strikes"), "summary": "s"}
    gone = {"id": "g1", "clause_key": "force_majeure", "change_type": "remove", "old": c.items[1].lettered[0], "new": "", "summary": "s"}
    for order in ([gone, modify], [modify, gone]):
        got, failed = points(order)
        assert failed == [] and got[1].startswith("2) ") and "any strikes" in got[1] and got[1].endswith("; or") and len(got) == 6, order
    last = {"id": "m7", "clause_key": "force_majeure", "change_type": "modify", "old": "7) epidemic.", "new": "7) epidemic and pandemic.", "summary": "s"}
    got, failed = points([last, {"id": "x", "clause_key": "force_majeure", "change_type": "remove", "old": c.items[1].lettered[5], "new": "", "summary": "s"}])
    assert failed == [] and got[-1] == "6) epidemic and pandemic." and got[-2].endswith("; or")


def test_a_point_is_found_by_what_it_says_not_by_its_number_or_ending():
    assert split_point("6) any change in law; or") == ("6)", "any change in law", "; or")
    assert split_point("(a) x; and")[:2] == ("(a)", "x") and split_point("(a) x; and")[2].replace(" ", "") == ";and"
    assert core_text("6) any change in law or government action; or") == core_text("5) Any change in law or government action.") == "any change in law or government action"
    assert core_text("plain sentence.") == "plain sentence" and core_text("") == ""


def test_two_points_that_say_the_same_are_not_guessed_between():
    library = library_for({})
    c = clause("force_majeure", library)
    c.items[1].lettered.append("9) epidemic.")
    failed = []
    apply_edits(library, [{"id": "d", "clause_key": "force_majeure", "change_type": "remove", "old": "8) epidemic.", "new": "", "summary": "s"}], failed)
    assert len(failed) == 1, "two points say 'epidemic': the recorded number is not there any more, so nothing is guessed"


def test_a_change_that_rewrote_a_whole_sentence_still_needs_that_sentence():
    LD_OLD, LD_NEW = old_ld(), new_ld()
    first = {"id": "a", "clause_key": "completion_deadline", "change_type": "modify", "old": LD_OLD, "new": LD_NEW, "summary": "s"}
    second = {"id": "b", "clause_key": "completion_deadline", "change_type": "modify", "old": LD_NEW, "new": LD_NEW.replace("per day", "each day"), "summary": "s"}
    failed = []
    library_for({}, edits=[second], failed=failed)
    assert [f["id"] for f in failed] == ["b"], "built on the first change, it cannot stand without it"
    failed = []
    library_for({}, edits=[first, second], failed=failed)
    assert failed == []


# ---------------------------------------------------------------- applying
def edit(change_type, old, new, key="completion_deadline"):
    return {"id": f"{key}-{change_type}-{len(new)}", "clause_key": key, "change_type": change_type, "old": old, "new": new, "summary": "s"}


def test_a_modify_replaces_the_sentence_in_place():
    lib = library_for({}, edits=[edit("modify", old_ld(), new_ld())])
    c = clause("completion_deadline", lib)
    assert [loc[0] for loc, t in units(c) if "0.03 (zero point zero three)" in t] == ["item"]
    assert not any("0.05 (zero point zero five)" in t for _, t in units(c))
    assert len(c.items) == len(clause("completion_deadline").items)


def test_the_caps_and_the_term_sheet_are_written_in_before_a_change_is_applied():
    lib = library_for({}, {"ld_cap": 3.0}, edits=[edit("modify", old_ld().replace("5 (five) percent", "3 (three) percent"), "X rewritten sentence.")])
    assert any(t == "X rewritten sentence." for _, t in units(clause("completion_deadline", lib)))     # the change matched the capped text


def test_a_table_cell_and_a_paragraph_can_be_changed():
    cell = "Payable upon completion of installation of the mounting system."
    lib = library_for({}, edits=[edit("modify", cell, "Payable upon completion of the mounting system.", "price_payment")])
    price = clause("price_payment", lib)
    assert price.tables[0]["rows"][1][3] == "Payable upon completion of the mounting system."
    assert price.table_rows[1][3] == "Payable upon completion of the mounting system."
    para = clause("notice").plain_paragraphs[1]
    lib = library_for({}, edits=[edit("modify", para, "Changed paragraph.", "notice")])
    assert clause("notice", lib).plain_paragraphs[1] == "Changed paragraph."


def test_an_added_sentence_goes_at_the_end_and_is_numbered():
    sentence = "The warranty period is 24 (twenty-four) months."
    lib = library_for({}, edits=[edit("add", "", sentence)])
    c = clause("completion_deadline", lib)
    assert c.items[-1].text == sentence and len(c.items) == len(clause("completion_deadline").items) + 1
    lines = next(items for n, cl, items in renumber_and_render(lib) if cl.key == "completion_deadline")
    assert lines[-1].endswith(sentence) and lines[-1].split(" ", 1)[0].startswith("6.")
    lib = library_for({}, edits=[edit("add", "", "A new last paragraph.", "notice")])
    assert clause("notice", lib).plain_paragraphs[-1] == "A new last paragraph."


def test_a_change_whose_sentence_is_gone_is_not_applied_and_is_reported():
    failed = []
    stale = edit("modify", "A sentence that is not in the clause.", "Whatever.")
    lib = library_for({}, edits=[stale, edit("modify", old_ld(), new_ld()), edit("add", "", "x y.", "no_such_clause")], failed=failed)
    assert [e["id"] for e in failed] == [stale["id"], failed[1]["id"]] and len(failed) == 2
    assert any("0.03" in t for _, t in units(clause("completion_deadline", lib)))                       # the good one still went in


def test_changes_are_applied_in_order_and_a_second_change_builds_on_the_first():
    first = edit("modify", old_ld(), new_ld())
    second = edit("modify", new_ld(), new_ld().replace("per day", "each day"))
    lib = library_for({}, edits=[first, second])
    assert any("each day" in t for _, t in units(clause("completion_deadline", lib)))
    failed = []
    library_for({}, edits=[second, first], failed=failed)                    # the other way round, the second has nothing to change yet
    assert len(failed) == 1


def test_a_fresh_library_has_none_of_the_changes():
    library_for({}, edits=[edit("modify", old_ld(), new_ld())])
    assert any("0.05 (zero point zero five)" in t for _, t in units(clause("completion_deadline")))


def test_the_changed_sentence_reaches_the_word_file():
    lib = library_for({}, edits=[edit("modify", old_ld(), new_ld()), edit("add", "", "The warranty period is 24 (twenty-four) months.")])
    buf = io.BytesIO()
    render_docx(lib, buf, header=build_header(None))
    buf.seek(0)
    text = "\n".join(p.text for p in Document(buf).paragraphs)
    assert "0.03 (zero point zero three) percent of the Contract Price per day" in text and "0.05 (zero point zero five)" not in text
    assert "The warranty period is 24 (twenty-four) months." in text


# ---------------------------------------------------------------- before and after
def test_both_sides_are_whole_and_bold_what_changes():
    old = "The Owner may claim damages at the rate of 0.05 (zero point zero five) percent per day."
    before, after = diff_markup(old, old.replace("0.05 (zero point zero five)", "0.03 (zero point zero three)"))
    assert "~~" not in before, "nothing is struck through"
    assert "**0.05**" in before and "**five)**" in before and "**0.03**" in after and "**three)**" in after, "the changed words are bold on both sides"
    assert before.replace("**", "").replace("\\", "") == old and after.replace("**", "") == old.replace("0.05", "0.03").replace("five", "three")
    assert before.startswith("The Owner may claim") and after.startswith("The Owner may claim"), "both sides are the whole text"
    assert diff_markup("same words", "same words") == ("same words", "same words")
    assert diff_markup("pay within 30 days", "pay within 30 days and 15 hours")[0] == "pay within 30 days"           # words only put in: nothing bold in the current wording
    assert diff_markup("a [b] *c*", "a [b] *d*") == ("a \\[b\\] **\\*c\\***", "a \\[b\\] **\\*d\\***")           # markdown characters are escaped


def test_added_sentence_does_not_repeat_the_clause():
    c = clause("suspension")
    own = " ".join(text for _, text in units(c))
    new = "The Contractor may suspend the Works if any installment remains unpaid more than 15 days after its due date."
    assert without_existing(c, own + " " + new) == new, "the clause's own sentences the AI wrote again are taken out"
    assert without_existing(c, new) == new, "a new sentence alone is not touched"
    assert without_existing(c, own) == "", "nothing new is left when the AI wrote only the clause again"
    assert without_existing(c, new + " " + own) == new, "also when the new sentence comes first"


def test_norm_is_one_line_with_single_spaces():
    assert norm("  a \n b\t c ") == "a b c" and norm(None) == ""


if __name__ == "__main__":
    tests = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    for name, fn in tests:
        fn()
        print(f"PASS {name}")
    print(f"\n{len(tests)} wording tests passed.")
