"""Tests for reading the term sheet Excel by its labels. Run: python -m tests.test_xlsx_labels"""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl

from core.xlsx_parser import FIELD_LABELS, _label_key, parse_xlsx, parse_xlsx_report
from tests.xlsx_helper import workbook

TEMPLATE = Path(__file__).resolve().parent.parent / "data" / "term_sheet_template_EN.xlsx"
VALUES = {"project_name": "Rooftop PV", "project_location": "Rayong", "capacity": "0.5 MW", "owner_name": "Owner Co.",
          "contract_price_number": "1000000", "vat_treatment": "exclusive", "owner_contact_email": "o@x.example",
          "contractor_contact_phone": "03-333-4444", "completion_date": "10 October 2026"}


def test_every_label_is_read_wherever_its_row_is():
    ts, rep = parse_xlsx_report(workbook(VALUES))
    assert rep.missing == [] and rep.conflicts == []                              # a label with an empty value is not missing
    for name, value in VALUES.items():
        assert getattr(ts, name) == ("0.5" if name == "capacity" else value), name


def test_rows_inserted_or_reordered_change_nothing():
    base = parse_xlsx(workbook(VALUES))
    reordered = parse_xlsx(workbook(VALUES, order=list(reversed(list(FIELD_LABELS))), start_row=40))
    assert [getattr(reordered, f) for f in VALUES] == [getattr(base, f) for f in VALUES]
    wb = openpyxl.load_workbook(workbook(VALUES)); ws = wb.active
    ws.insert_rows(5, amount=3); ws.insert_rows(20)                               # rows inserted in the middle
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    ts, rep = parse_xlsx_report(buf)
    assert rep.missing == [] and [getattr(ts, f) for f in VALUES] == [getattr(base, f) for f in VALUES]


def test_a_deleted_row_is_reported_and_the_others_still_read_correctly():
    wb = openpyxl.load_workbook(workbook(VALUES)); ws = wb.active
    row = next(r[0].row for r in ws.iter_rows(min_col=2, max_col=2) if r[0].value == "VAT treatment")
    ws.delete_rows(row)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    ts, rep = parse_xlsx_report(buf)
    assert rep.missing == ["VAT treatment"] and ts.vat_treatment == ""
    assert (ts.project_name, ts.owner_name, ts.contract_price_number, ts.completion_date) == ("Rooftop PV", "Owner Co.", "1000000", "10 October 2026")


def test_a_label_that_is_gone_is_reported_not_guessed():
    wb = openpyxl.load_workbook(workbook(VALUES)); ws = wb.active
    for row in ws.iter_rows(min_col=2, max_col=2):
        if row[0].value == "VAT treatment":
            row[0].value = "VAT handling"                                          # renamed
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    ts, rep = parse_xlsx_report(buf)
    assert ts.vat_treatment == "" and "VAT treatment" in rep.missing and ts.project_name == "Rooftop PV"


def test_case_spaces_and_a_colon_in_a_label_do_not_matter():
    wb = openpyxl.load_workbook(workbook(VALUES)); ws = wb.active
    for row in ws.iter_rows(min_col=2, max_col=2):
        if row[0].value == "Project name":
            row[0].value = "  PROJECT   name: "
        if row[0].value == "Capacity (kWp or MW)":
            row[0].value = "capacity (kWp or MW)\u00a0"
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    ts = parse_xlsx(buf)
    assert ts.project_name == "Rooftop PV" and ts.capacity == "0.5" and ts.capacity_unit == "MW"


def test_a_repeated_label_uses_the_first_filled_copy_and_reports_a_difference():
    def two(first, second):
        return parse_xlsx_report(workbook({"capacity": first}, extra_rows=[("Capacity (kWp or MW)", second, None)]))
    ts, rep = two("100 MW", "100 MW")
    assert rep.conflicts == [] and ts.capacity == "100"
    ts, rep = two("100MW", "100 mw ")                                              # only spacing and case differ
    assert rep.conflicts == []
    ts, rep = two("100 MW", "250 MW")
    assert ts.capacity == "100" and len(rep.conflicts) == 1 and "appears 2 times" in rep.conflicts[0] and "250" not in ts.capacity
    ts, rep = two(None, "250 MW")                                                  # the first copy is empty: the filled one counts
    assert ts.capacity == "250" and rep.conflicts == []


def test_a_cell_with_several_lines_keeps_its_lines():
    ts = parse_xlsx(workbook({"owner_signatory": "Nat\r\n  Somchai  \n\nMalee", "project_location": "Grattan Street\nParkville"},
                             extra_rows=[("Installment 1 (60%)", "On signing:\n(1) the Agreement\n(2) the purchase order", "(1) Drawings\n(2) Shop drawings"),
                                         ("Installment 2 (40 %)", "On handover", None)]))
    assert ts.owner_signatory == "Nat\nSomchai\n\nMalee" and ts.project_location == "Grattan Street\nParkville"
    assert [(i.number, i.percent) for i in ts.installments] == [("1", "60"), ("2", "40")]
    assert ts.installments[0].condition == "On signing:\n(1) the Agreement\n(2) the purchase order"
    assert ts.installments[0].required_docs == "(1) Drawings\n(2) Shop drawings"


def test_a_value_continued_on_the_row_below_is_not_read():
    wb = openpyxl.load_workbook(workbook(VALUES)); ws = wb.active
    ws["C2"] = "continued here, on a row with no label"                            # row 2 has no label (the fields start at row 3)
    ws.cell(row=ws.max_row + 1, column=3, value="more text with no label")
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    ts = parse_xlsx(buf)
    assert "continued here" not in repr(ts) and "more text" not in repr(ts)


def test_a_file_without_the_labels_is_recognised_as_not_the_template():
    wb = openpyxl.Workbook(); wb.active["B2"], wb.active["C2"] = "Something else", "x"
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    ts, rep = parse_xlsx_report(buf)
    assert rep.nothing_found and len(rep.missing) == len(FIELD_LABELS)


def test_only_the_first_sheet_is_read_and_installment_rows_are_found_by_name():
    wb = openpyxl.load_workbook(workbook({"project_name": "First sheet"}, extra_rows=[("Installment 2 (50%)", "B", None), ("Installment 1 (50%)", "A", None)]))
    second = wb.create_sheet("Other"); second["B1"], second["C1"] = "Project name", "Second sheet"
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    ts = parse_xlsx(buf)
    assert ts.project_name == "First sheet" and [i.condition for i in ts.installments] == ["A", "B"]


def test_the_template_in_the_repository_has_every_label_once_and_reads_cleanly():
    assert TEMPLATE.exists(), "data/term_sheet_template_EN.xlsx is missing"
    ws = openpyxl.load_workbook(TEMPLATE).active
    labels = [_label_key(r[0].value) for r in ws.iter_rows(min_col=2, max_col=2) if isinstance(r[0].value, str) and r[0].value.strip()]
    assert sorted({l for l in labels if labels.count(l) > 1}) == [], "a label appears more than once in the template"
    ts, rep = parse_xlsx_report(TEMPLATE)
    assert rep.missing == [] and rep.conflicts == [] and len(ts.installments) >= 1


def test_the_template_has_a_governing_law_row_that_is_read_into_the_field():
    from core.xlsx_parser import FIELD_LABELS, parse_xlsx_report
    assert FIELD_LABELS["governing_law"] == "Governing law"
    ts, report = parse_xlsx_report(open("data/term_sheet_template_EN.xlsx", "rb"))
    assert ts.governing_law == "Thailand" and "Governing law" not in report.missing


if __name__ == "__main__":
    tests = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    for name, fn in tests:
        fn()
        print(f"PASS {name}")
    print(f"\n{len(tests)} Excel label tests passed.")
