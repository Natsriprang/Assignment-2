"""Tests for the capacity unit and the contract header. Run: python -m tests.test_header"""
import io
import sys
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
from docx import Document

from core.capacity import capacity_in, parse_capacity
from core.clause_data import build_library
from core.draft_state import library_for
from core.header import TITLE, build_header
from core.render import render_docx
from core.term_sheet_fields import TermSheet
from core.xlsx_parser import parse_xlsx


def full_term_sheet(**overrides):
    ts = TermSheet(project_name="Rooftop PV Project", project_location="Bangkok, Thailand", capacity="500",
                   contract_no="C-2026-001", contract_date="1 November 2026",
                   owner_name="Solar Owner Co., Ltd.", owner_reg_no="0105500000001", owner_address="1 Owner Road, Bangkok",
                   contractor_name="Sun Build Co., Ltd.", contractor_reg_no="0105500000002", contractor_address="2 Builder Road, Bangkok")
    for k, v in overrides.items():
        setattr(ts, k, v)
    return ts


# ---------------------------------------------------------------------------- capacity
def test_parse_capacity_reads_numbers_and_units():
    cases = {"500": (500, "kWp"), "500 kWp": (500, "kWp"), "500kWp": (500, "kWp"), "0.5 MW": (Decimal("0.5"), "MW"),
             "1.25MW": (Decimal("1.25"), "MW"), "1,250": (1250, "kWp"), "500 kW": (500, "kWp"), "0.5 MWp": (Decimal("0.5"), "MW"),
             "  500  ": (500, "kWp"), "500.00": (500, "kWp")}
    for text, (value, unit) in cases.items():
        got = parse_capacity(text)
        assert got[0] == Decimal(value) and got[1] == unit, (text, got)


def test_parse_capacity_rejects_what_is_not_a_positive_number():
    for text in ("", "abc", "0", "-5", "5 GW", "five hundred", "NaN", "Infinity", None):
        assert parse_capacity(text)[0] is None, text


def test_the_default_unit_applies_only_when_the_text_has_none():
    assert parse_capacity("0.5", "MW") == (Decimal("0.5"), "MW")
    assert parse_capacity("500", "MW") == (Decimal("500"), "MW")
    assert parse_capacity("500 kWp", "MW") == (Decimal("500"), "kWp")      # the unit written in the text wins
    assert parse_capacity("500", "bogus")[1] == "kWp"


def test_conversion_is_exact_in_both_directions():
    kwp = lambda cap, unit="kWp": capacity_in(SimpleNamespace(capacity=cap, capacity_unit=unit), "kWp")
    mw = lambda cap, unit="kWp": capacity_in(SimpleNamespace(capacity=cap, capacity_unit=unit), "MW")
    assert mw("500") == "0.5" and mw("1250") == "1.25" and mw("333") == "0.333" and mw("500.00") == "0.5"
    assert kwp("0.5", "MW") == "500" and kwp("1.25", "MW") == "1250" and kwp("500") == "500" and mw("0.5", "MW") == "0.5"
    assert kwp("abc") is None and mw("") is None


def test_an_older_term_sheet_object_without_the_unit_is_read_as_kwp():
    assert capacity_in(SimpleNamespace(capacity="500"), "MW") == "0.5"
    assert capacity_in(None, "MW") is None and capacity_in(SimpleNamespace(), "kWp") is None


def test_xlsx_import_splits_the_unit_from_the_number():
    def sheet(value):
        wb = openpyxl.Workbook()
        wb.active["C6"] = value
        buf = io.BytesIO(); wb.save(buf); buf.seek(0)
        return parse_xlsx(buf)
    ts = sheet("0.5 MW"); assert (ts.capacity, ts.capacity_unit) == ("0.5", "MW")
    ts = sheet("1,250 kWp"); assert (ts.capacity, ts.capacity_unit) == ("1250", "kWp")
    ts = sheet(500); assert (ts.capacity, ts.capacity_unit) == ("500", "kWp")
    ts = sheet("to be confirmed"); assert (ts.capacity, ts.capacity_unit) == ("to be confirmed", "kWp")    # not guessed or lost


# ---------------------------------------------------------------------------- header
def test_a_complete_term_sheet_fills_every_header_field_it_has_data_for():
    lines = build_header(full_term_sheet())
    assert lines == [
        ("title", "Construction and Installation Service Agreement"),
        ("centre", "for the Photovoltaic Solar Power Rooftop Facility"),
        ("centre", "Rooftop PV Project with 0.5 MW"),
        ("centre", "Contract no.: C-2026-001"),
        ("para", "This construction and installation service agreement (the \u201cAgreement\u201d) is made on 1 November 2026 "
                 "(the \u201cdate of this Agreement\u201d) at [*]."),
        ("para", "Between:"),
        ("bullet", "Solar Owner Co., Ltd., a company incorporated under the laws of [*] having the company registration number "
                   "0105500000001, whose registered head office is at 1 Owner Road, Bangkok, hereinafter referred to as the "
                   "\u201cOwner\u201d and;"),
        ("bullet", "Sun Build Co., Ltd., a company incorporated under the laws of Thailand having the company registration number "
                   "0105500000002, whose registered head office is at 2 Builder Road, Bangkok, hereinafter referred to as the "
                   "\u201cContractor\u201d"),
        ("para", "The Owner and the Contractor are hereinafter referred to as the \u201cParties\u201d"),
        ("para", "WHEREAS"),
        ("indent", "The Owner wishes to install the photovoltaic solar power system on the rooftop (the \u201cSystem\u201d) located at "
                   "Bangkok, Thailand (the \u201cProject\u201d). The Contractor is an expert in providing services on design, construct "
                   "and installation of the System."),
        ("indent", "The Parties wish to enter into the subcontracting services on construction and installation agreement for the "
                   "photovoltaic solar power rooftop in accordance with the terms and conditions set out in this Agreement."),
    ]


def test_missing_values_stay_visible_as_placeholders_and_nothing_is_invented():
    lines = build_header(TermSheet())
    texts = [t for _, t in lines]
    assert "[*] with [*] MW" in texts and "Contract no.: [*]" in texts
    assert sum(t.count("[*]") for t in texts) == 13           # the template has 13 blanks, none added or lost
    assert "laws of Thailand" in texts[7] and "laws of [*]" in texts[6]          # the Contractor is Thai, the Owner's country is not known
    assert not any("{" in t or "None" in t for t in texts)


def test_place_of_signing_and_the_owners_country_are_never_guessed():
    texts = [t for _, t in build_header(full_term_sheet())]
    assert texts[4].endswith("at [*].") and "laws of [*] having" in texts[6]


def test_capacity_in_the_header_is_in_megawatts_whatever_unit_was_entered():
    for cap, unit, expect in (("500", "kWp", "0.5"), ("1250", "kWp", "1.25"), ("0.5", "MW", "0.5"), ("2", "MW", "2")):
        assert build_header(full_term_sheet(capacity=cap, capacity_unit=unit))[2][1] == f"Rooftop PV Project with {expect} MW"
    assert build_header(full_term_sheet(capacity="tbc"))[2][1] == "Rooftop PV Project with [*] MW"


def test_header_accepts_none_and_old_term_sheet_objects():
    assert build_header(None)[0] == ("title", TITLE)
    assert build_header(SimpleNamespace(project_name="P", capacity="500"))[2][1] == "P with 0.5 MW"


# ---------------------------------------------------------------------------- Word file
def docx_lines(header):
    library = library_for({})
    buf = io.BytesIO()
    render_docx(library, buf, header=header)
    buf.seek(0)
    return [p.text for p in Document(buf).paragraphs if p.text.strip()]


def test_the_word_file_starts_with_the_header_and_clause_numbering_is_unchanged():
    lines = docx_lines(build_header(full_term_sheet()))
    assert lines[0] == TITLE and lines[1] == "for the Photovoltaic Solar Power Rooftop Facility"
    assert "Rooftop PV Project with 0.5 MW" in lines and "Contract no.: C-2026-001" in lines and "WHEREAS" in lines
    assert any(l.startswith("\u2022 Solar Owner Co., Ltd.") for l in lines) and any(l.startswith("\u2022 Sun Build Co., Ltd.") for l in lines)
    first_clause = next(i for i, l in enumerate(lines) if l.startswith("1.\t"))
    assert lines.index("WHEREAS") < first_clause
    assert lines[first_clause] == "1.\tConstruction and Installation Service Agreement"


def test_without_a_header_the_old_title_only_output_is_unchanged():
    lines = docx_lines(None)
    assert lines[0] == TITLE and lines[1].startswith("1.\t")


if __name__ == "__main__":
    tests = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    for name, fn in tests:
        fn()
        print(f"PASS {name}")
    print(f"\n{len(tests)} capacity and header tests passed.")
