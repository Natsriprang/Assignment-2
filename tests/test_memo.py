"""Tests for the internal approval memo (core/memo.py, core/memo_render.py). Run: python -m tests.test_memo"""
import io
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document

from core.header import PLACEHOLDER
from core.memo import build_memo
from core.memo_render import render_memo_docx
from core.term_sheet_fields import TermSheet
from tests.test_interpret import sample_term_sheet

TS = sample_term_sheet()
DAY = date(2026, 10, 9)


def word_text(memo):
    buf = io.BytesIO()
    render_memo_docx(memo, buf)
    buf.seek(0)
    doc = Document(buf)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


def test_fields_follow_the_form():
    memo = build_memo(TS, today=DAY)
    assert [label for label, _ in memo["fields"]] == ["To", "Subject", "Documents proposed for signing", "Urgency"]
    assert memo["date"] == "9 October 2026"


def test_subject_and_document_name_the_deal():
    fields = dict(build_memo(TS, today=DAY)["fields"])
    for text in (fields["Subject"], fields["Documents proposed for signing"]):
        assert TS.project_name in text and "Construction and Installation Service Agreement" in text
    assert TS.owner_name in fields["Subject"] and TS.contractor_name in fields["Subject"]
    assert "1 copy (original)" in fields["Documents proposed for signing"]


def test_to_and_urgency_are_left_for_the_person():
    fields = dict(build_memo(TS, today=DAY)["fields"])
    assert fields["To"] == PLACEHOLDER and fields["Urgency"] == PLACEHOLDER


def test_two_sign_off_boxes_proposer_and_legal():
    boxes = build_memo(TS, today=DAY)["boxes"]
    assert [b["heading"] for b in boxes] == ["Proposed by", "Reviewed by the Legal Department"]
    assert boxes[1]["line"] == "Legal Department"
    assert "executive" not in word_text(build_memo(TS, today=DAY)).lower()


def test_key_terms_are_the_facts_table():
    memo = build_memo(TS, today=DAY)
    labels = [label for label, _ in memo["key_terms"]]
    assert "Contract price" in labels and "Parties" in labels
    assert "Contract price: THB 10,000,000" in word_text(memo)


def test_empty_term_sheet_shows_placeholders_and_does_not_fail():
    memo = build_memo(TermSheet(), today=DAY)
    assert PLACEHOLDER in dict(memo["fields"])["Subject"]
    assert "Approval" in word_text(memo)
    assert build_memo(None, today=DAY)["fields"][1][1].count(PLACEHOLDER) >= 2


def test_word_file_has_the_form_and_the_note_that_it_is_internal():
    text = word_text(build_memo(TS, today=DAY))
    for expected in ("Date: 9 October 2026", "Details", "Submitted for your consideration.", "Proposed by", "Reviewed by the Legal Department",
                     "not part of the contract"):
        assert expected in text, expected


def test_memo_is_pure_code_no_ai_client_needed():
    import inspect
    import core.memo, core.memo_render
    for module in (core.memo, core.memo_render):
        assert "client" not in inspect.getsource(module).lower()


if __name__ == "__main__":
    tests = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    for name, fn in tests:
        fn()
        print("PASS", name)
    print(f"{len(tests)} memo tests passed.")
