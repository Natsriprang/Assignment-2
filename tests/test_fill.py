"""Tests for writing term sheet values into the clauses, the blanks list, and the new term sheet fields.
Run: python -m tests.test_fill"""
import sys
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from core.blanks import find_blanks
from core.clause_data import build_library
from core.clause_params import int_words
from core.draft_state import library_for
from core.fill import DEADLINE_BODY, DEADLINE_TAIL, deadline_sentence, fill_deadline, fill_slots, fill_template, fill_warnings, format_date, looks_like_date, parse_decimal
from core.header import build_header
from core.renumber import renumber_and_render
from core.term_sheet_fields import Installment, TermSheet, ensure_fields, signing_place


def full(**over):
    ts = TermSheet(project_name="Rooftop PV", project_location="Bangkok", capacity="500", contract_no="C-1", contract_date="1 Nov 2026",
                   owner_name="Solar Owner Co., Ltd.", owner_country="Singapore", owner_reg_no="O-1", owner_address="1 Owner Road, Bangkok",
                   contractor_name="Sun Build Co., Ltd.", contractor_reg_no="C-2", contractor_address="2 Builder Road, Bangkok",
                   contract_price_number="10000000", contract_price_words="ten million baht", vat_treatment="exclusive",
                   attachment_no="3", deadline_mode="period", completion_days="120",
                   signing_place_choice="owner", owner_signatory="Nat Sriprang", contractor_signatory="Noom Sriprang",
                   owner_contact_name="John Smith", owner_contact_position="Director", owner_contact_location="1 Owner Road, Bangkok",
                   owner_contact_phone="02-111-2222", owner_contact_email="john@owner.example",
                   contractor_contact_name="Jan Lee", contractor_contact_position="Engineer", contractor_contact_location="2 Builder Road, Bangkok",
                   contractor_contact_phone="03-333-4444", contractor_contact_email="jan@builder.example")
    ts.installments = [Installment(str(i + 1), p) for i, p in enumerate(["10", "30", "30", "20", "10"])]
    for k, v in over.items():
        setattr(ts, k, v)
    return ts


def clause(library, key):
    return next(c for c in library if c.key == key)


def lines(library, key):
    return next(items for n, c, items in renumber_and_render(library) if c.key == key)


# ---------------------------------------------------------------------------- no term sheet, no change
def test_nothing_is_written_without_a_term_sheet():
    plain = [(c.plain_paragraphs, [i.text for i in c.items], c.tables) for c in build_library()]
    filled = [(c.plain_paragraphs, [i.text for i in c.items], c.tables) for c in fill_template(build_library(), None)]
    assert plain == filled


def test_a_blank_term_sheet_changes_only_the_defaults_it_holds():
    lib = library_for({}, None, None, TermSheet())
    assert "within 30 (thirty) working days" in lines(lib, "price_payment")[2]          # payment days default to 30
    assert "THB [*] ([*])" in lines(lib, "price_payment")[0] and "installed capacity of [*] kWp" in clause(lib, "construction_terms").plain_paragraphs[0]


# ---------------------------------------------------------------------------- clause by clause
def test_capacity_and_location_in_clause_one():
    text = lambda **o: clause(library_for({}, None, None, full(**o)), "construction_terms").plain_paragraphs[0]
    assert "installed capacity of 500 kWp (\u201cWorks\u201d) at Bangkok in accordance" in text()
    assert "installed capacity of 0.5 MW (\u201cWorks\u201d)" in text(capacity="0.5", capacity_unit="MW")
    assert "installed capacity of [*] MW (\u201cWorks\u201d) at Bangkok" in text(capacity="", capacity_unit="MW")          # the unit still follows
    assert "(\u201cWorks\u201d) at [*] in accordance" in text(project_location="")


def test_attachment_number():
    para = lambda **o: clause(library_for({}, None, None, full(**o)), "attachments").plain_paragraphs[1]
    assert "this Agreement and Attachment 3, the provisions" in para()
    assert "Attachment [*], the provisions" in para(attachment_no="")


def test_price_and_its_words():
    first = lambda **o: lines(library_for({}, None, None, full(**o)), "price_payment")[0]
    assert "THB 10,000,000 (ten million baht), exclusive of VAT, the terms" in first()
    assert "THB 10,000,000.50 (" in first(contract_price_number="10000000.50")
    assert "THB 1,250,000 (" in first(contract_price_number="1,250,000")
    assert "THB 10,000,000 ([*]), exclusive" in first(contract_price_words="")
    assert "THB ten million ([*])" in first(contract_price_number="ten million", contract_price_words="")       # typed text is kept as typed


def test_payment_days_in_figures_and_words():
    days = lambda **o: lines(library_for({}, None, None, full(**o)), "price_payment")[2]
    assert "within 30 (thirty) working days" in days()
    assert "within 45 (forty-five) working days" in days(payment_days="45")
    assert "within 120 (one hundred and twenty) working days" in days(payment_days="120")
    for bad in ("abc", "0", "", "1000", "-5", "2.5"):
        assert "within [*] ([*]) working days" in days(payment_days=bad), bad


def test_installment_table_rows_in_order_with_amounts():
    rows = lambda **o: clause(library_for({}, None, None, full(**o)), "price_payment").tables[0]["rows"]
    r = rows()
    assert [row[:3] for row in r] == [["1", "10%", "1,000,000.00"], ["2", "30%", "3,000,000.00"], ["3", "30%", "3,000,000.00"],
                                      ["4", "20%", "2,000,000.00"], ["5", "10%", "1,000,000.00"]]
    assert r[0][3].startswith("Payable upon the Owner's confirmation")                               # the template's conditions stay
    assert rows(contract_price_number="")[0][:3] == ["1", "10%", "[*]"]                               # no price, no amount
    ts = full(); ts.installments[0].percent = "12.5 %"
    assert clause(library_for({}, None, None, ts), "price_payment").tables[0]["rows"][0][1:3] == ["12.5%", "1,250,000.00"]


def test_fewer_or_more_installments_than_the_table_has_rows():
    ts = full(); ts.installments = ts.installments[:3]
    rows = clause(library_for({}, None, None, ts), "price_payment").tables[0]["rows"]
    assert rows[2][:3] == ["3", "30%", "3,000,000.00"] and rows[3][:3] == ["[*]", "[*]", "[*]"] and rows[4][:3] == ["[*]", "[*]", "[*]"]
    assert any("3 installments" in w and "4-5 are left blank" in w for w in fill_warnings(full(installments=full().installments[:3])))


def test_the_term_sheets_conditions_replace_the_templates_and_a_blank_keeps_them():
    ts = full()
    ts.installments[0].condition = "Upon signing of this Agreement"
    ts.installments[2].condition = "  Upon delivery of the PV modules to site  "
    rows = clause(library_for({}, None, None, ts), "price_payment").tables[0]["rows"]
    assert rows[0][3] == "Upon signing of this Agreement" and rows[2][3] == "Upon delivery of the PV modules to site"
    assert rows[1][3] == "Payable upon completion of installation of the mounting system."      # blank: the template's
    assert rows[3][3] == "Payable upon completion of installation of the inverter."


def test_a_condition_with_several_lines_keeps_its_lines():
    ts = full()
    ts.installments[1].condition = "On completion and the Owner's acceptance of:\r\n(1) walkway   installation\n\n  (2) life line installation  "
    ts.installments += [Installment("6", "5", condition="Final:\n(1) handover"), Installment("7", "5", condition="Done")]
    price = clause(library_for({}, None, None, ts), "price_payment")
    rows = price.tables[0]["rows"]
    assert rows[1][3] == "On completion and the Owner's acceptance of:\n(1) walkway installation\n(2) life line installation"
    assert rows[5][3] == "Final:\n(1) handover" and rows[6][3] == "Done"


def test_the_lines_of_a_condition_reach_the_word_file_as_line_breaks():
    import io
    from docx import Document
    from core.render import render_docx
    ts = full()
    ts.installments[1].condition = "Accept:\n(1) walkway\n(2) life line\n(3) mounting"
    buf = io.BytesIO()
    render_docx(library_for({}, None, None, ts), buf, header=build_header(ts)); buf.seek(0)
    pay = [t for t in Document(buf).tables if t.rows[0].cells[0].text == "Installment"][0]
    cell = pay.rows[2].cells[3]
    assert cell._tc.xml.count("<w:br/>") == 3 and cell.text == "Accept:\n(1) walkway\n(2) life line\n(3) mounting"
    assert pay.rows[1].cells[3]._tc.xml.count("<w:br/>") == 0                      # a one-line cell has none


def test_every_installment_of_the_term_sheet_gets_a_row():
    ts = full(contract_price_number="1000000")
    ts.installments = [Installment(str(i + 1), p) for i, p in enumerate(["10", "20", "20", "15", "15", "10", "10"])]
    ts.installments[5].condition, ts.installments[6].condition = "Upon final acceptance", ""
    lib = library_for({}, None, None, ts)
    price = clause(lib, "price_payment")
    rows = price.tables[0]["rows"]
    assert len(rows) == 7 and [r[0] for r in rows] == [str(i) for i in range(1, 8)]
    assert rows[5] == ["6", "10%", "100,000.00", "Upon final acceptance"] and rows[6] == ["7", "10%", "100,000.00", "[*]"]
    assert rows[0][3].startswith("Payable upon the Owner's confirmation")                        # a blank condition keeps the template's
    assert [tuple(r) for r in rows] == price.table_rows
    assert fill_warnings(ts) == []                                                               # nothing is left out, so nothing to warn about
    again = fill_template(lib, ts)                                                               # filling twice does not add rows
    assert len(clause(again, "price_payment").tables[0]["rows"]) == 7
    left = find_blanks(build_header(ts), renumber_and_render(lib))
    assert [b.where.endswith("table, row 7)") for b in left if "table" in b.where] == [True]     # only the missing condition is blank


def test_the_percentages_and_amounts_check_covers_every_installment():
    ts = full(contract_price_number="1000")
    ts.installments = [Installment(str(i + 1), "14.3") for i in range(7)] + []
    assert not any("round" in w for w in fill_warnings(ts))                                       # 100.1 %, not 100: that is the percent warning's job
    ts.installments = [Installment(str(i + 1), p) for i, p in enumerate(["33.33", "33.33", "33.34"])]
    assert not any("round" in w for w in fill_warnings(ts))


def test_the_deadline_sentence_constants_are_the_templates():
    template = next(c for c in build_library() if c.key == "completion_deadline").items[0].text
    assert template == DEADLINE_BODY + DEADLINE_TAIL                     # if the template is edited, this fails instead of the fill going quiet


def test_the_deadline_is_one_way_or_the_other_never_both():
    first = lambda **o: lines(library_for({}, None, None, full(**o)), "completion_deadline")[0].split(" ", 1)[1]       # without "6.1 "
    dates = first(deadline_mode="calendar", start_date="03 October 2026", completion_date="10 October 2026")
    assert dates == ("The Owner agrees that the Contractor shall commence the Works under this Agreement on or before 3 October 2026 "
                     "and the Contractor agrees to complete the Works on or before 10 October 2026.")
    assert first() == "The Contractor agrees to complete the Works under this Agreement within 120 days from the signing date of this Agreement."
    for text in (dates, first()):
        assert "or within" not in text and "150" not in text            # no fallback alternative, in either mode


def test_a_period_has_only_a_completion():
    first = lambda **o: lines(library_for({}, None, None, full(**o)), "completion_deadline")[0]
    assert "commence" not in first() and "commence" not in first(start_date="1 Nov 2026", completion_date="1 Dec 2026")
    assert "on or before 3 October 2026 and" in first(deadline_mode="calendar", start_date="2026-10-03 00:00:00", completion_date="x")
    assert "on or before 03/10/2026" in first(deadline_mode="calendar", start_date="03/10/2026", completion_date="x")      # kept as typed
    assert "on or before 3 ตุลาคม 2569" in first(deadline_mode="calendar", start_date="3 ตุลาคม 2569", completion_date="x")  # so is any other form


def test_only_what_the_mode_names_is_used():
    first = lambda **o: lines(library_for({}, None, None, full(**o)), "completion_deadline")[0]
    stale = dict(start_date="1 Nov 2026", completion_date="1 Dec 2026")                          # typed earlier, then the mode was switched
    assert "within 120 days" in first(**stale) and "1 Nov 2026" not in first(**stale) and "1 Dec 2026" not in first(**stale)
    cal = dict(deadline_mode="calendar", start_date="1 Nov 2026", completion_date="1 Dec 2026", completion_days="30")
    assert "on or before 1 December 2026." in first(**cal) and "30 days" not in first(**cal)


def test_without_a_mode_or_a_value_the_blank_stays():
    first = lambda **o: lines(library_for({}, None, None, full(**o)), "completion_deadline")[0].split(" ", 1)[1]
    assert first(deadline_mode="") == DEADLINE_BODY + DEADLINE_TAIL                    # nothing chosen: the template as it is
    for bad in ("abc", "0", "1000", "-5", "2.5", "", "120 days"):
        assert first(completion_days=bad) == "The Contractor agrees to complete the Works under this Agreement within [*] days from the signing date of this Agreement.", bad
    assert "on or before [*] and" in first(deadline_mode="calendar", start_date="", completion_date="1 Dec 2026")
    assert deadline_sentence("some other sentence", full()) == "some other sentence"                         # an edited template is left alone
    assert deadline_sentence(DEADLINE_BODY + DEADLINE_TAIL, full(deadline_mode="nonsense")) == DEADLINE_BODY + DEADLINE_TAIL


def test_a_sentence_with_a_different_number_of_blanks_is_left_alone():
    assert fill_deadline("no blanks here", ["x", "y"]) == "no blanks here"
    assert fill_slots("one [*] and [*]", ["a"]) == "one [*] and [*]"
    assert fill_slots("one [*] and [*]", ["a", None]) == "one a and [*]"
    assert fill_slots("no blank", []) == "no blank"


def test_only_the_intended_clauses_change():
    before = {c.key: ([*c.plain_paragraphs], [i.text for i in c.items]) for c in library_for({})}
    after = {c.key: ([*c.plain_paragraphs], [i.text for i in c.items]) for c in library_for({}, None, None, full())}
    assert sorted(k for k in before if before[k] != after[k]) == ["attachments", "completion_deadline", "construction_terms", "price_payment"]
    tables = lambda lib: {c.key: [r[:] for t in c.tables for r in t["rows"]] for c in lib}
    assert sorted(k for k, v in tables(library_for({}, None, None, full())).items() if v != tables(library_for({})).get(k)) == ["notice", "price_payment"]


def test_the_vat_wording_is_still_resolved():
    assert "exclusive of VAT" in lines(library_for({}, None, None, full()), "price_payment")[0]
    assert "inclusive / exclusive of VAT" in lines(library_for({}, None, None, full(vat_treatment="")), "price_payment")[0]


# ---------------------------------------------------------------------------- warnings
def test_a_complete_term_sheet_has_no_warnings():
    assert fill_warnings(full()) == [] and fill_warnings(None) == []


def test_warnings_name_what_was_left_blank_and_why():
    w = fill_warnings(full(payment_days="abc"))
    assert any("Payment days" in x and "abc" in x for x in w)
    assert any("empty" in x for x in fill_warnings(full(signing_place_choice="contractor", contractor_address="")))


# ---------------------------------------------------------------------------- blanks list
def blanks_for(ts):
    lib = library_for({}, None, None, ts)
    return find_blanks(build_header(ts), renumber_and_render(lib))


def test_a_complete_term_sheet_leaves_no_blanks_except_what_it_does_not_hold():
    left = blanks_for(full())
    assert left == [], [(b.where, b.context) for b in left]


def test_a_blank_term_sheet_lists_every_blank_with_where_it_is():
    left = blanks_for(TermSheet())
    wheres = {b.where for b in left}
    assert "Contract header" in wheres and any("Contract Price" in w for w in wheres) and any("table, row 1" in w for w in wheres)
    assert any("installed capacity of [*] kWp" in b.context for b in left)


def test_a_blank_inside_a_drafted_clause_is_listed():
    lib = library_for({}, None, [{"key": "drafted_1", "title": "Warranty", "sentences": ["A penalty of [*]% per day."], "before": "end"}], full())
    left = find_blanks(build_header(full()), renumber_and_render(lib))
    assert [(b.where.split(". ", 1)[1], "[*]%" in b.context) for b in left] == [("Warranty", True)]


# ---------------------------------------------------------------------------- the new fields
def test_signing_place_is_read_from_the_chosen_address():
    assert signing_place(full()) == "1 Owner Road, Bangkok"
    assert signing_place(full(signing_place_choice="contractor")) == "2 Builder Road, Bangkok"
    assert signing_place(full(signing_place_choice="other", signing_place_other=" Chiang Mai ")) == "Chiang Mai"
    assert signing_place(full(signing_place_choice="")) == "" and signing_place(SimpleNamespace()) == ""
    assert signing_place(full(owner_address="")) == ""


def test_the_header_uses_the_countries_and_the_place():
    texts = [t for _, t in build_header(full())]
    assert texts[4].endswith("at 1 Owner Road, Bangkok.")
    assert "incorporated under the laws of Singapore having" in texts[6]
    assert "incorporated under the laws of Thailand having" in texts[7]
    assert "laws of [*] having" in [t for _, t in build_header(full(owner_country=""))][6]
    assert "laws of Thailand having" in [t for _, t in build_header(SimpleNamespace(project_name="P"))][7]            # older object


def test_the_defaults_and_old_objects():
    ts = TermSheet()
    assert ts.payment_days == "30" and ts.contractor_country == "Thailand" and ts.owner_country == "" and ts.signing_place_choice == ""
    old = ensure_fields(SimpleNamespace(project_name="P"))
    assert (old.payment_days, old.contractor_country, old.owner_country, old.attachment_no) == ("30", "Thailand", "", "")
    assert int_words(30) == "thirty" and parse_decimal("1,000.50") and parse_decimal("0") is None and parse_decimal("x") is None


if __name__ == "__main__":
    tests = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    for name, fn in tests:
        fn()
        print(f"PASS {name}")
    print(f"\n{len(tests)} fill and blanks tests passed.")
