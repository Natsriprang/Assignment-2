"""Builds small term sheet workbooks for the tests: a label in column B and its value in column C, one field per row."""
import io

import openpyxl

from core.xlsx_parser import FIELD_LABELS


def workbook(values=None, *, extra_rows=(), start_row=3, order=None):
    """`values` maps a term sheet field name to the cell value. `extra_rows` are (label, value, note) rows written after
    them (installments, section headers). `order` lists field names to write in that order (default: FIELD_LABELS order)."""
    values = values or {}
    wb = openpyxl.Workbook()
    ws = wb.active
    row = start_row
    for name in (order or list(FIELD_LABELS)):
        ws.cell(row=row, column=2, value=FIELD_LABELS[name])
        if name in values:
            ws.cell(row=row, column=3, value=values[name])
        row += 1
    for label, value, note in extra_rows:
        ws.cell(row=row, column=2, value=label)
        ws.cell(row=row, column=3, value=value)
        ws.cell(row=row, column=4, value=note)
        row += 1
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
