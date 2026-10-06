"""Tests for the notice table, the signature block, and where they appear (preview data, Word file, blanks).
Run: python -m tests.test_notice_signature"""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document

from core.blanks import find_blanks
from core.clause_data import build_library
from core.draft_state import library_for
from core.field_rules import split_names
from core.header import build_header
from core.render import render_docx
from core.renumber import renumber_and_render
from core.signatures import CLOSING, SIGNATURE_LINE, build_signature_block, names_of, signature_texts
from core.term_sheet_fields import TermSheet
from tests.test_fill import clause, full


def notice_rows(ts):
    return clause(library_for({}, None, None, ts), "notice").tables[0]["rows"]


# ---------------------------------------------------------------- the notice table
def test_the_template_has_a_notice_table_with_a_block_per_party():
    rows = clause(build_library(), "notice").tables[0]["rows"]
    assert clause(build_library(), "notice").tables[0]["headers"] == ["Party", "Item", "Details"]
    assert [r[:2] for r in rows] == [["Owner", "Company name"], ["", "Attention"], ["", "Address"], ["", "Phone"], ["", "Email"],
                                     ["Contractor", "Company name"], ["", "Attention"], ["", "Address"], ["", "Phone"], ["", "Email"]]
    assert all(r[2] == "[*]" for r in rows)


def test_the_table_is_filled_from_the_party_names_and_the_contacts():
    rows = notice_rows(full())
    assert [r[2] for r in rows] == ["Solar Owner Co., Ltd.", "John Smith, Director", "1 Owner Road, Bangkok", "02-111-2222", "john@owner.example",
                                    "Sun Build Co., Ltd.", "Jan Lee, Engineer", "2 Builder Road, Bangkok", "03-333-4444", "jan@builder.example"]
    assert [r[0] for r in rows] == ["Owner", "", "", "", "", "Contractor", "", "", "", ""]               # the party names stay as the template has them


def test_a_missing_contact_value_stays_a_blank_and_a_lone_position_is_kept():
    rows = notice_rows(full(owner_contact_phone="", owner_contact_email="  ", contractor_contact_name=""))
    assert rows[3][2] == "[*]" and rows[4][2] == "[*]" and rows[6][2] == "Engineer"
    assert notice_rows(full(owner_contact_name="", owner_contact_position=""))[1][2] == "[*]"
    assert all(r[2] == "[*]" for r in notice_rows(TermSheet())[1:5]) and notice_rows(None) is not None


def test_a_contact_that_runs_over_several_lines_is_one_line_in_the_table():
    assert notice_rows(full(owner_contact_location="1 Owner Road\nBangkok 10110"))[2][2] == "1 Owner Road Bangkok 10110"


def test_filling_twice_changes_nothing_more():
    from core.fill import fill_template
    lib = library_for({}, None, None, full())
    once = [r[:] for r in clause(lib, "notice").tables[0]["rows"]]
    assert [r for r in clause(fill_template(lib, full(owner_contact_phone="999")), "notice").tables[0]["rows"]] == once


# ---------------------------------------------------------------- the signature block
def test_each_party_has_its_company_and_a_signing_line_per_signatory():
    block = build_signature_block(full())
    assert block["closing"] == CLOSING and [p["role"] for p in block["parties"]] == ["Owner", "Contractor"]
    assert block["parties"][0]["lines"] == ["OWNER", "Solar Owner Co., Ltd.", "", SIGNATURE_LINE, "Nat Sriprang", "Authorised signatory",
                                            "", SIGNATURE_LINE, "Ploy Wong", "Witness"]
    assert block["parties"][1]["lines"] == ["CONTRACTOR", "Sun Build Co., Ltd.", "", SIGNATURE_LINE, "Noom Sriprang", "Authorised signatory",
                                            "", SIGNATURE_LINE, "Mali Chan", "Witness"]


def test_several_signatories_each_get_a_line_and_a_missing_name_is_a_blank():
    lines = build_signature_block(full(owner_signatory="Nat Sriprang\nSomchai Jaidee"))["parties"][0]["lines"]
    assert lines[2:] == ["", SIGNATURE_LINE, "Nat Sriprang", "Authorised signatory", "", SIGNATURE_LINE, "Somchai Jaidee", "Authorised signatory",
                         "", SIGNATURE_LINE, "Ploy Wong", "Witness"]
    assert split_names("A B; C D, E F & G H and I J\nK L") == ["A B", "C D", "E F", "G H", "I J", "K L"] and split_names("  ") == []
    blank = build_signature_block(full(owner_signatory="", contractor_name=""))
    assert blank["parties"][0]["lines"][4] == "[*]" and blank["parties"][1]["lines"][1] == "[*]"
    assert build_signature_block(None)["parties"][0]["lines"][1] == "[*]"


def test_witnesses_sign_under_the_authorised_signatories_of_their_own_party_and_more_can_be_added():
    owner = build_signature_block(full(owner_signatory="Nat Sriprang\nSomchai Jaidee", owner_witnesses="Ploy Wong\nKanya Rit"))["parties"][0]["lines"]
    titles = [x for x in owner if x in ("Authorised signatory", "Witness")]
    assert titles == ["Authorised signatory", "Authorised signatory", "Witness", "Witness"], "all the signatories first, then the witnesses"
    assert owner[2:] == ["", SIGNATURE_LINE, "Nat Sriprang", "Authorised signatory", "", SIGNATURE_LINE, "Somchai Jaidee", "Authorised signatory",
                         "", SIGNATURE_LINE, "Ploy Wong", "Witness", "", SIGNATURE_LINE, "Kanya Rit", "Witness"]
    contractor = build_signature_block(full(owner_witnesses="Ploy Wong\nKanya Rit"))["parties"][1]["lines"]
    assert contractor.count("Witness") == 1 and "Ploy Wong" not in contractor, "each party has its own witnesses"


def test_a_party_without_a_witness_name_still_has_one_witness_line_with_a_blank():
    for empty in ("", "  \n "):
        lines = build_signature_block(full(owner_witnesses=empty))["parties"][0]["lines"]
        assert lines[-4:] == ["", SIGNATURE_LINE, "[*]", "Witness"]
    assert build_signature_block(None)["parties"][1]["lines"][-4:] == ["", SIGNATURE_LINE, "[*]", "Witness"]
    ts = full(owner_witnesses="")
    found = find_blanks(build_header(ts), renumber_and_render(library_for({}, None, None, ts)), build_signature_block(ts))
    assert [b.where for b in found] == ["Signature block"], "the missing witness name is listed as a blank"


def test_one_name_per_line_when_there_are_line_breaks_and_the_old_separators_when_there_are_not():
    assert names_of("Somchai Jaidee, Dr.\nNat Sriprang") == ["Somchai Jaidee, Dr.", "Nat Sriprang"], "a comma in a name is kept when names are on separate lines"
    assert names_of("Nat Sriprang; Somchai Jaidee") == ["Nat Sriprang", "Somchai Jaidee"]
    assert names_of("") == [] and names_of(None) == [] and names_of("\n \n") == []


# ---------------------------------------------------------------- blanks
def test_blanks_in_the_notice_table_and_the_signature_block_are_listed_where_they_are():
    ts = full(owner_contact_email="", contractor_signatory="")
    lib = library_for({}, None, None, ts)
    found = find_blanks(build_header(ts), renumber_and_render(lib), build_signature_block(ts))
    wheres = [b.where for b in found]
    assert any("Notice (table, row 5)" in w for w in wheres) and "Signature block" in wheres
    assert any(b.where == "Signature block" and "[*]" in b.context for b in found)
    assert find_blanks(build_header(full()), renumber_and_render(library_for({}, None, None, full())), build_signature_block(full())) == []
    assert not [b for b in find_blanks(build_header(ts), renumber_and_render(lib)) if b.where == "Signature block"]      # no block given, none searched


# ---------------------------------------------------------------- the Word file
def make_docx(ts, with_block=True):
    buf = io.BytesIO()
    render_docx(library_for({}, None, None, ts), buf, header=build_header(ts), signature=build_signature_block(ts) if with_block else None)
    buf.seek(0)
    return Document(buf)


def test_the_word_file_has_the_notice_table_after_its_first_paragraph():
    d = make_docx(full())
    body = [(type(c).__name__, c) for c in d.element.body.iterchildren()]
    texts = [p.text for p in d.paragraphs]
    i = next(n for n, t in enumerate(texts) if t.startswith("For the purpose of the performance of this Agreement being effective"))
    # the element right after that paragraph is the table, and the second paragraph follows the table
    first = d.paragraphs[i]._p
    after = first.getnext()
    assert after.tag.endswith("}tbl")
    table = [t for t in d.tables if t.rows[0].cells[0].text == "Party"][0]
    assert [c.text for c in table.rows[1].cells] == ["Owner", "Company name", "Solar Owner Co., Ltd."]
    assert [c.text for c in table.rows[10].cells] == ["", "Email", "jan@builder.example"]
    nxt = after.getnext()
    while nxt is not None and nxt.tag.endswith("}p") and not "".join(nxt.itertext()).strip():
        nxt = nxt.getnext()                                           # the empty paragraph after any table
    assert "".join(nxt.itertext()).startswith("Any communication with the addressee")


def test_the_word_file_ends_with_the_signature_block():
    d = make_docx(full())
    assert d.paragraphs[-1].text == CLOSING or any(p.text == CLOSING for p in d.paragraphs)
    last = d.tables[-1]
    assert len(last.rows) == 1 and len(last.columns) == 2
    owner, contractor = last.rows[0].cells
    assert [p.text for p in owner.paragraphs] == ["OWNER", "Solar Owner Co., Ltd.", "", SIGNATURE_LINE, "Nat Sriprang", "Authorised signatory",
                                                  "", SIGNATURE_LINE, "Ploy Wong", "Witness"]
    assert [p.text for p in contractor.paragraphs][-2:] == ["Mali Chan", "Witness"], "the witnesses come last in the column"
    assert owner.paragraphs[0].runs[0].bold and not owner.paragraphs[1].runs[0].bold
    order = list(d.element.body.iterchildren())
    closing = next(pp for pp in d.paragraphs if pp.text == CLOSING)
    last_clause = next(pp for pp in d.paragraphs if "governed by and interpreted in accordance with the laws of Thailand" in pp.text)
    assert order.index(closing._p) > order.index(last_clause._p), "the block comes after the last clause"
    assert order.index(last.rows[0].cells[0]._tc.getparent().getparent()) == order.index(closing._p) + 1, "the table follows the closing sentence"


def test_without_a_block_the_word_file_is_as_before():
    d = make_docx(full(), with_block=False)
    assert not any(p.text == CLOSING for p in d.paragraphs) and all(t.rows[0].cells[0].text != "OWNER" for t in d.tables)


def test_every_blank_of_a_complete_draft_is_gone_in_the_word_file():
    text = "\n".join(p.text for p in make_docx(full()).paragraphs) + "\n".join(c.text for t in make_docx(full()).tables for r in t.rows for c in r.cells)
    assert "[*]" not in text


if __name__ == "__main__":
    tests = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    for name, fn in tests:
        fn()
        print(f"PASS {name}")
    print(f"\n{len(tests)} notice and signature tests passed.")
