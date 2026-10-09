"""Tests for what is written in bold (core/emphasis.py) and how the Word file uses it. Run: python -m tests.test_emphasis"""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document

from core.draft_state import library_for
from core.emphasis import party_names, segments
from core.header import build_header
from core.render import render_docx
from tests.test_header import full_term_sheet


def test_defined_terms_in_quotation_marks_are_bold_with_their_marks():
    assert segments("made on 1 May (the “date of this Agreement”) at X") == [
        ("made on 1 May (the ", False), ("“date of this Agreement”", True), (") at X", False)]


def test_text_without_terms_is_one_plain_piece_and_pieces_join_back():
    assert segments("plain text") == [("plain text", False)]
    text = "the “Owner” and the “Contractor”, ACME Ltd."
    assert "".join(p for p, _ in segments(text, ["ACME Ltd."])) == text


def test_names_are_bold_only_when_given():
    assert ("ACME Ltd.", True) in segments("ACME Ltd. is here", ["ACME Ltd."])
    assert ("ACME Ltd.", True) not in segments("ACME Ltd. is here")


def test_party_names_are_names_and_registration_numbers_longest_first():
    ts = full_term_sheet()
    names = party_names(ts)
    assert ts.owner_name in names and ts.owner_reg_no in names and ts.contractor_name in names and ts.contractor_reg_no in names
    assert names == sorted(names, key=len, reverse=True)
    assert party_names(None) == []


def bold_pieces(paragraph):
    return [r.text for r in paragraph.runs if r.bold]


def test_word_file_bolds_party_names_numbers_and_defined_terms_in_the_header():
    ts = full_term_sheet()
    buf = io.BytesIO()
    render_docx(library_for({}), buf, header=build_header(ts), bold_names=party_names(ts))
    buf.seek(0)
    paras = Document(buf).paragraphs
    owner = next(p for p in paras if p.text.startswith("• " + ts.owner_name))
    assert ts.owner_name in bold_pieces(owner) and ts.owner_reg_no in bold_pieces(owner) and "“Owner”" in bold_pieces(owner)
    assert ts.owner_address not in " ".join(bold_pieces(owner))
    first = next(p for p in paras if p.text.startswith("This construction"))
    assert bold_pieces(first) == ["“Agreement”", "“date of this Agreement”"]


def test_word_file_bolds_defined_terms_in_the_clauses_and_names_only_in_the_header():
    ts = full_term_sheet()
    buf = io.BytesIO()
    render_docx(library_for({}), buf, header=build_header(ts), bold_names=party_names(ts))
    buf.seek(0)
    paras = Document(buf).paragraphs
    found = {t for p in paras for t in bold_pieces(p)}
    assert "“Force Majeure”" in found and "“Works”" in found
    body = [p for p in paras if p.text.startswith(("1.1", "2.1"))]
    assert all(ts.owner_name not in bold_pieces(p) for p in body)


def test_signature_block_bolds_the_role_and_the_company_name():
    ts = full_term_sheet()
    from core.signatures import build_signature_block
    buf = io.BytesIO()
    render_docx(library_for({}), buf, header=build_header(ts), signature=build_signature_block(ts))
    buf.seek(0)
    cells = Document(buf).tables[-1].rows[0].cells
    for cell, company in zip(cells, (ts.owner_name, ts.contractor_name)):
        bold = [r.text for p in cell.paragraphs for r in p.runs if r.bold]
        assert company in bold and len(bold) == 2
        assert "Authorised signatory" not in bold


if __name__ == "__main__":
    tests = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    for name, fn in tests:
        fn()
        print("PASS", name)
    print(f"{len(tests)} emphasis tests passed.")
