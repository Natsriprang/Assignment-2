"""Write the approval memo (core.memo.build_memo) to a Word file: a table of the fields, the details, and the two sign-off boxes."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def _cell(cell, text, bold=False):
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold


def keep_row_together(row):
    """The sign-off boxes must not be split across two pages."""
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:cantSplit"))


def render_memo_docx(memo, out_path):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.space_after = Pt(4)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(memo["title"])
    run.bold = True
    run.font.size = Pt(13)
    date_line = doc.add_paragraph()
    date_line.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    date_line.add_run(f"Date: {memo['date']}").bold = True

    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    for label, text in memo["fields"]:
        left, right = table.add_row().cells
        _cell(left, label, bold=True)
        _cell(right, text)
        left.width, right.width = Cm(4.5), Cm(12.0)
    table.columns[0].width, table.columns[1].width = Cm(4.5), Cm(12.0)
    doc.add_paragraph()

    doc.add_paragraph().add_run(memo["details_label"]).bold = True
    for text in memo["details"]:
        doc.add_paragraph(text)
    if memo["key_terms"]:
        doc.add_paragraph(memo["key_terms_label"])
        for label, value in memo["key_terms"]:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            p.add_run("• ")
            p.add_run(f"{label}: ").bold = True
            p.add_run(value)
    doc.add_paragraph(memo["closing"]).paragraph_format.space_before = Pt(12)

    boxes = doc.add_table(rows=1, cols=len(memo["boxes"]))
    boxes.style = "Table Grid"
    keep_row_together(boxes.rows[0])
    for cell, box in zip(boxes.rows[0].cells, memo["boxes"]):
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].add_run(box["heading"]).bold = True
        for _ in range(3):                                     # room to sign
            cell.add_paragraph()
        for text in (box["name"], box["line"]):
            p = cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(text)
    foot = doc.add_paragraph()
    foot.paragraph_format.space_before = Pt(10)
    foot.add_run(memo["footer"]).italic = True
    doc.save(out_path)
