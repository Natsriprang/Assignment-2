"""Render a clause list to a .docx."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from core.emphasis import segments
from core.renumber import renumber_and_render


def _add_text(paragraph, text, names=(), bold=False):
    """Write `text` into the paragraph; defined terms in quotation marks and `names` are in bold."""
    for piece, strong in segments(text, names):
        paragraph.add_run(piece).bold = bold or strong or None


def _add_table(doc, table_data):
    t = doc.add_table(rows=1, cols=len(table_data["headers"]))
    t.style = "Light Grid Accent 1"
    for cell, text in zip(t.rows[0].cells, table_data["headers"]):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True
    for row in table_data["rows"]:
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    doc.add_paragraph()


def _add_signature_block(doc, block):
    """The closing sentence, then one borderless two-column table: a column per party, a paragraph per line."""
    doc.add_paragraph()
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _add_text(para, block["closing"])
    table = doc.add_table(rows=1, cols=len(block["parties"]))
    for cell, party in zip(table.rows[0].cells, block["parties"]):
        for n, line in enumerate(party["lines"]):
            p = cell.paragraphs[0] if n == 0 else cell.add_paragraph()
            run = p.add_run(line)
            run.bold = n < 2                                    # the role is the heading of the column, and the company's name is in bold


def render_docx(clauses, out_path, header=None, signature=None, bold_names=()):
    """`header` is a list of (style, text) lines from core.header.build_header; without it only the title is written.
    `signature` is a block from core.signatures.build_signature_block, written after the last clause."""
    rendered = renumber_and_render(clauses)

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    recital = 0
    for style, text in header or [("title", "Construction and Installation Service Agreement")]:
        p = doc.add_paragraph()
        if style in ("title", "centre"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if style == "right":
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if style in ("para", "bullet", "numbered", "indent"):
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if style == "numbered":                                # a recital: a number in front, the text hanging
            recital += 1
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.first_line_indent = Cm(-0.7)
            text = f"{recital}.\t{text}"
        if style == "bullet":
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            text = "\u2022 " + text
        if style == "indent":
            p.paragraph_format.left_indent = Cm(1.0)
        if style == "title":
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(14)
        else:
            _add_text(p, text, bold_names if style == "bullet" else ())
    doc.add_paragraph()

    for number, clause, item_lines in rendered:
        h = doc.add_paragraph()
        r = h.add_run(f"{number}.\t{clause.title}")
        r.bold = True

        for n, p in enumerate(clause.plain_paragraphs):
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _add_text(para, p)
            if n == 0 and not item_lines:                      # a clause without numbered items: its table follows its first paragraph
                for table_data in clause.tables:
                    _add_table(doc, table_data)

        for i, line in enumerate(item_lines):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if line.startswith("    "):
                p.paragraph_format.left_indent = Cm(1.0)
                _add_text(p, line.strip())
            else:
                p.paragraph_format.left_indent = Cm(0.5)
                _add_text(p, line)
            if i == 0 and clause.tables:
                for table_data in clause.tables:
                    _add_table(doc, table_data)

        if not item_lines and not clause.plain_paragraphs:
            for table_data in clause.tables:
                _add_table(doc, table_data)

    if signature:
        _add_signature_block(doc, signature)

    doc.save(out_path)
    return rendered